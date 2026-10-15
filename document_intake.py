"""Bounded local extraction. Documents remain untrusted inputs to agents."""
import io, zipfile, re
from pypdf import PdfReader
from openpyxl import load_workbook
import domain as d

def extract(body,mime):
    if mime=='text/plain':return body.decode('utf-8-sig',errors='replace')[:30000]
    if mime=='application/pdf':
        reader=PdfReader(io.BytesIO(body))
        if reader.is_encrypted:d.fail('Unlock the PDF before extraction')
        if len(reader.pages)>100:d.fail('Maximum 100 PDF pages for intake')
        text='\n'.join(f'PAGE {i+1}\n{page.extract_text() or ""}' for i,page in enumerate(reader.pages[:30]))[:30000]
        if len(text.strip())<20:d.fail('This PDF needs OCR; upload a text PDF or transcribed source')
        return text
    if mime=='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
        with zipfile.ZipFile(io.BytesIO(body)) as z:
            if sum(x.file_size for x in z.infolist())>20*1024*1024:d.fail('Expanded workbook is too large')
        book=load_workbook(io.BytesIO(body),read_only=True,data_only=True);parts=[]
        try:
            for sheet in book.worksheets[:8]:
                parts.append('SHEET '+sheet.title)
                for row in sheet.iter_rows(max_row=500,max_col=30,values_only=True):parts.append(' | '.join(str(x) for x in row if x is not None))
        finally:book.close()
        return '\n'.join(parts)[:30000]
    d.fail('Image OCR is not configured. Upload a text PDF, XLSX or plain text source.',503)

def validate_proposal(proposal,source,kind):
    if not proposal:d.fail('Agent did not extract a document',502)
    fields=proposal['fields'];sources=proposal['sources']
    missing=list(proposal.get('missing',[]))
    for name in list(fields):
        if fields[name] is None:fields.pop(name);sources.pop(name,None);missing.append(name)
    allowed={'quotation':{'supplier','quantity','unit_price','freight','discount','tax_rate','currency','delivery_date','quality','reliability','risk','certifications','payment_terms','minimum_order','supplier_email','payee_reference'},'invoice':{'reference','quantity','unit_price','total','currency','date','due_date','quote_id'},'receipt':{'quantity','reference','date','quote_id'},'inspection':{'receipt_id','accepted','reason'}}[kind]
    if set(fields)-allowed:d.fail('Document extraction contains unsupported fields',502)
    for name,value in list(fields.items()):
        excerpt=sources.get(name)
        # An empty optional value is absence of evidence, not a certified fact.
        if name in {'certifications','supplier_email','payee_reference'} and value in ('',[],None) and not excerpt:
            fields.pop(name);sources.pop(name,None);missing.append(name);continue
        if isinstance(excerpt,str) and excerpt.strip() and excerpt not in source:
            # Correct harmless case/spacing/punctuation changes by finding the REAL
            # contiguous source span. Never store an LLM paraphrase as a quote.
            tokens=re.findall(r'\d+(?:\.\d+)?|[^\W\d_]+',excerpt,flags=re.UNICODE)
            if tokens:
                match=re.search(r'\b'+r'[\W_]*'.join(re.escape(t) for t in tokens)+r'\b',source,flags=re.IGNORECASE)
                if match:excerpt=source[match.start():match.end()];sources[name]=excerpt
        if not isinstance(excerpt,str) or not excerpt.strip() or excerpt not in source:d.fail('Extracted field lacks an exact source excerpt: '+name,502)
        if name in {'quantity','unit_price','total','freight','discount','tax_rate','quality','reliability','risk','minimum_order','accepted'}:
            expected=d.dec(value)
            numbers=re.findall(r'(?<![\w.])-?\d[\d,]*(?:\.\d+)?(?![\w.])',excerpt)
            if not any(d.dec(n.replace(',',''))==expected for n in numbers if not n.startswith('-')):d.fail('Extracted number is not supported by its source: '+name,502)
    # A model may list fields as missing while also returning validated values.
    missing=set(missing)-set(fields)
    return {'fields':fields,'sources':sources,'missing':sorted(missing),'interpretation':'Groq extraction; source-linked interpretation, subject to business validation and human approvals'}
